from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/"docs/10_user_guide/InkTrace-小白写手完整操作手册.md"
OUTPUT=ROOT/"docs/10_user_guide/InkTrace-小白写手完整操作手册.docx"

def set_font(style,name,size,color="000000",bold=False):
    style.font.name=name; style.font.size=Pt(size); style.font.color.rgb=RGBColor.from_string(color); style.font.bold=bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")

def page_number(paragraph):
    run=paragraph.add_run(); fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); run._r.addnext(fld)

doc=Document(); section=doc.sections[0]
section.page_width=Inches(8.5); section.page_height=Inches(11)
section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
section.header_distance=section.footer_distance=Inches(0.492)

normal=doc.styles["Normal"]; set_font(normal,"Calibri",11); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.25
for key,size,before,after,color in [("Title",28,0,10,"203748"),("Subtitle",13,0,18,"667580"),("Heading 1",16,18,10,"2E74B5"),("Heading 2",13,14,7,"2E74B5"),("Heading 3",12,10,5,"1F4D78")]:
    style=doc.styles[key]; set_font(style,"Calibri",size,color,key!="Subtitle"); style.paragraph_format.space_before=Pt(before); style.paragraph_format.space_after=Pt(after); style.paragraph_format.keep_with_next=True
for key in ("List Bullet","List Number"):
    style=doc.styles[key]; set_font(style,"Calibri",11); style.paragraph_format.left_indent=Inches(0.375); style.paragraph_format.first_line_indent=Inches(-0.188); style.paragraph_format.space_after=Pt(4); style.paragraph_format.line_spacing=1.25

header=section.header.paragraphs[0]; header.text="InkTrace · 小白写手操作手册"; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.style,"Calibri",9,"667580")
footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT; footer.add_run("第 "); page_number(footer); footer.add_run(" 页")

lines=SOURCE.read_text(encoding="utf-8").splitlines()
toc_items=[line.strip()[3:] for line in lines if line.startswith("## ")]
for raw in lines:
    line=raw.strip()
    if not line or line=="---": continue
    if line.startswith("# "):
        p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(line[2:])
        continue
    if line.startswith("版本："):
        p=doc.add_paragraph(style="Subtitle"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(line.replace("  ","")); continue
    if line.startswith("适用对象："):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(line); r.italic=True; r.font.color.rgb=RGBColor.from_string("667580")
        doc.add_page_break()
        doc.add_paragraph("快速目录",style="Heading 1")
        for item in toc_items:
            doc.add_paragraph(item,style="List Bullet")
        doc.add_page_break(); continue
    if line.startswith("#### "): doc.add_paragraph(line[5:],style="Heading 3"); continue
    if line.startswith("### "): doc.add_paragraph(line[4:],style="Heading 2"); continue
    if line.startswith("## "): doc.add_paragraph(line[3:],style="Heading 1"); continue
    if line.startswith("> "):
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.25); p.paragraph_format.right_indent=Inches(0.25); p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(10)
        r=p.add_run(re.sub(r"\*\*","",line[2:])); r.bold=True; r.font.color.rgb=RGBColor.from_string("1F4D78")
        pPr=p._p.get_or_add_pPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"F4F6F9"); pPr.append(shd); continue
    if re.match(r"^\d+\. ",line): doc.add_paragraph(re.sub(r"^\d+\. ","",line),style="List Number"); continue
    if line.startswith("- [ ] "):
        p=doc.add_paragraph(style="List Bullet"); p.add_run("☐ "+line[6:]); continue
    if line.startswith("- "): doc.add_paragraph(line[2:],style="List Bullet"); continue
    p=doc.add_paragraph()
    parts=re.split(r"(\*\*.*?\*\*)",line)
    for part in parts:
        if not part: continue
        if part.startswith("**") and part.endswith("**"): p.add_run(part[2:-2]).bold=True
        else: p.add_run(part.replace("  ",""))

doc.core_properties.title="InkTrace 小白写手完整操作手册"; doc.core_properties.subject="InkTrace V2.0 小说写作操作指南"
doc.save(OUTPUT)
print(OUTPUT)
